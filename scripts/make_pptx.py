"""Generates PRESENTATION.pptx -- 15 slides, for hackathon submission.

    cd /d E:\\Crossover_point && conda activate crossover && ^
        python scripts\\make_pptx.py
"""
from __future__ import annotations

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

NAVY = RGBColor(0x0F, 0x17, 0x2A)
BLUE = RGBColor(0x3B, 0x82, 0xF6)
GREEN = RGBColor(0x22, 0xC5, 0x5E)
AMBER = RGBColor(0xF5, 0x9E, 0x0B)
RED = RGBColor(0xEF, 0x44, 0x44)
GREY = RGBColor(0x6B, 0x72, 0x80)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
DARK = RGBColor(0x11, 0x18, 0x27)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]


def slide():
    return prs.slides.add_slide(BLANK)


def bg(s, color=WHITE):
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = color


def box(s, l, t, w, h):
    tb = s.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tb.text_frame.word_wrap = True
    return tb.text_frame


def para(tf, text, size=18, color=DARK, bold=False, align=PP_ALIGN.LEFT, first=False, space_after=8, italic=False):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment = align
    p.space_after = Pt(space_after)
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    r.font.color.rgb = color
    r.font.name = "Segoe UI"
    return p


def title_bar(s, kicker, title, color=BLUE):
    tf = box(s, 0.6, 0.35, 12.1, 0.4)
    para(tf, kicker.upper(), size=13, color=color, bold=True, first=True)
    tf2 = box(s, 0.6, 0.72, 12.1, 0.9)
    para(tf2, title, size=30, color=DARK, bold=True, first=True)
    ln = s.shapes.add_shape(1, Inches(0.6), Inches(1.55), Inches(12.1), Pt(2.2))
    ln.fill.solid(); ln.fill.fore_color.rgb = color; ln.line.fill.background()


def bullets(s, items, l=0.7, t=1.9, w=11.9, h=5.2, size=18):
    tf = box(s, l, t, w, h)
    for i, (txt, kw) in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(12)
        r = p.add_run(); r.text = "•  "; r.font.size = Pt(size); r.font.color.rgb = BLUE; r.font.bold = True
        r2 = p.add_run(); r2.text = txt; r2.font.size = Pt(size); r2.font.color.rgb = DARK; r2.font.name = "Segoe UI"
        if kw:
            r2.font.bold = False


def stat_card(s, l, t, w, h, value, label, color):
    box_shape = s.shapes.add_shape(1, Inches(l), Inches(t), Inches(w), Inches(h))
    box_shape.fill.solid(); box_shape.fill.fore_color.rgb = RGBColor(0xF3, 0xF4, 0xF6)
    box_shape.line.color.rgb = color; box_shape.line.width = Pt(1.5)
    tf = box_shape.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = value; r.font.size = Pt(30); r.font.bold = True; r.font.color.rgb = color
    p2 = tf.add_paragraph(); p2.alignment = PP_ALIGN.CENTER
    r2 = p2.add_run(); r2.text = label; r2.font.size = Pt(12); r2.font.color.rgb = GREY


def footer(s, n):
    tf = box(s, 12.3, 7.05, 0.9, 0.35)
    para(tf, f"{n} / 15", size=11, color=GREY, first=True, align=PP_ALIGN.RIGHT)


# ---------------------------------------------------------------- 1. Title
s = slide(); bg(s, NAVY)
tf = box(s, 1, 2.5, 11.3, 1.3)
para(tf, "AI CROSSOVER POINT", size=54, color=WHITE, bold=True, first=True, align=PP_ALIGN.CENTER)
tf2 = box(s, 1, 3.75, 11.3, 0.8)
para(tf2, "Making a 7B AI model answer faster by letting a tiny model guess first",
     size=20, color=RGBColor(0x93, 0xC5, 0xFD), first=True, align=PP_ALIGN.CENTER)
tf3 = box(s, 1, 6.4, 11.3, 0.6)
para(tf3, "A working prototype measuring speculative decoding on consumer hardware  |  Hackathon Submission",
     size=14, color=RGBColor(0xCB, 0xD5, 0xE1), first=True, align=PP_ALIGN.CENTER)

# ---------------------------------------------------------------- 2. The problem
s = slide(); bg(s)
title_bar(s, "The Problem", "Big AI models are accurate, but slow to talk to")
bullets(s, [
    ("Large language models (like the 7B model used here) generate text one word "
     "at a time — each word needs a full, expensive pass through the whole network.", None),
    ("This makes chatbots and coding assistants feel sluggish, especially on a "
     "normal consumer PC instead of a data-center server.", None),
    ("The obvious fix — using a smaller, faster model — sacrifices answer quality. "
     "We want the big model's quality at closer to the small model's speed.", None),
    ("Question this project asks: can a small 0.5B \"draft\" model guess ahead, and "
     "have the big 7B model just check the guesses, instead of writing every word itself?", None),
], size=19)
footer(s, 2)

# ---------------------------------------------------------------- 3. The idea
s = slide(); bg(s)
title_bar(s, "The Idea", "Speculative decoding, in one sentence", GREEN)
tf = box(s, 0.8, 2.0, 11.7, 1.3)
para(tf, "A small, fast model guesses the next few words. The big, accurate model "
         "checks all those guesses in one single pass, and keeps whichever ones it agrees with.",
     size=22, color=DARK, bold=True, first=True)
bullets(s, [
    ("Analogy: a junior writer drafts a paragraph quickly; a senior editor reads it "
     "once and approves the parts that are already correct — instead of writing every word themselves.", None),
    ("Because the big model has the final say on every single word, the output text "
     "is provably identical to what it would have written alone. Nothing about the "
     "quality changes — only how fast it arrives.", None),
    ("This technique is not new — it is published research (Leviathan et al. 2022, "
     "Chen et al. 2023). This project's contribution is measuring how well it "
     "actually works on ordinary consumer hardware, which nobody had done for this setup.", None),
], t=3.5, size=18)
footer(s, 3)

# ---------------------------------------------------------------- 4. How one round works
s = slide(); bg(s)
title_bar(s, "How It Works", "One round of speculative decoding")
steps = [
    ("1", "Draft guesses", "The small 0.5B model quickly proposes the next few tokens (words/word-pieces).", BLUE),
    ("2", "Target verifies", "The big 7B model checks all guesses at once, in a single forward pass.", GREEN),
    ("3", "Keep or correct", "Correct guesses are kept; the first wrong one is replaced with the big model's own choice, and a bonus token is sometimes free.", AMBER),
]
x = 0.7
for num, head, desc, color in steps:
    card = s.shapes.add_shape(1, Inches(x), Inches(2.1), Inches(3.9), Inches(4.3))
    card.fill.solid(); card.fill.fore_color.rgb = RGBColor(0xF9, 0xFA, 0xFB)
    card.line.color.rgb = color; card.line.width = Pt(1.5)
    tf = card.text_frame; tf.word_wrap = True
    tf.margin_left = Inches(0.25); tf.margin_right = Inches(0.25); tf.margin_top = Inches(0.3)
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = num; r.font.size = Pt(40); r.font.bold = True; r.font.color.rgb = color
    p2 = tf.add_paragraph(); p2.alignment = PP_ALIGN.CENTER
    r2 = p2.add_run(); r2.text = head; r2.font.size = Pt(20); r2.font.bold = True; r2.font.color.rgb = DARK
    p3 = tf.add_paragraph()
    r3 = p3.add_run(); r3.text = desc; r3.font.size = Pt(14); r3.font.color.rgb = RGBColor(0x37, 0x41, 0x51)
    x += 4.1
tf = box(s, 0.7, 6.55, 11.9, 0.7)
para(tf, "Result: one pass of the big model can produce several words instead of just one -- when guesses are good.",
     size=14, color=GREY, first=True, italic=True)
footer(s, 4)

# ---------------------------------------------------------------- 5. Why output never changes
s = slide(); bg(s)
title_bar(s, "Correctness Guarantee", "Why the output text never changes", GREEN)
bullets(s, [
    ("The big model (\"target\") independently computes its own top choice at every "
     "position during verification -- exactly as if it had generated that word alone.", None),
    ("A guessed word is accepted ONLY if it exactly matches what the target would "
     "have picked anyway. So every kept word is one the target model chose itself.", None),
    ("At the first disagreement, the guess is discarded and the target's own word is "
     "used instead -- the rest of that round's guesses are thrown away, unseen.", None),
    ("This means the small draft model never decides WHAT is said -- only HOW MANY "
     "words can be produced per expensive pass. Quality and behaviour are untouched.", None),
], size=19)
footer(s, 5)

# ---------------------------------------------------------------- 6. Test setup
s = slide(); bg(s)
title_bar(s, "Test Setup", "How this was measured", BLUE)
bullets(s, [
    ("Hardware: one consumer NVIDIA RTX 3060 (12 GB), an ordinary gaming GPU -- "
     "not a data-center card.", None),
    ("Target (big) model: Qwen2.5-Coder-7B-Instruct, compressed to 4-bit precision.", None),
    ("Draft (small) model: Qwen2.5-Coder-0.5B-Instruct, running at normal precision.", None),
    ("Both models share the exact same vocabulary, which is required for the "
     "guess-and-check comparison to make sense.", None),
    ("Test set: 210 prompts, hand-checked, split evenly across six real-world "
     "categories: structured data, code, math, reasoning, prose writing, and translation.", None),
], size=18)
footer(s, 6)

# ---------------------------------------------------------------- 7. Test cases proving correctness
s = slide(); bg(s)
title_bar(s, "Proof It Works", "Test cases that verify correctness", GREEN)
bullets(s, [
    ("2,520 total timed runs across two execution modes were checked byte-by-byte "
     "against plain, unmodified generation from the big model alone.", None),
    ("96.7% - 96.8% of all runs produced text IDENTICAL down to the exact character.", None),
    ("Every one of the small number of differences was individually re-examined. "
     "Every single one occurred exactly where the big model's top two possible "
     "words were tied in floating-point precision -- meaning even plain, ordinary "
     "generation would be ambiguous there too.", None),
    ("Zero decoder bugs were found. This was checked mathematically (ULP-level "
     "floating point analysis), not just by eye.", None),
    ("An automated pytest suite of 85 tests also passes, covering cache handling, "
     "the decoding logic, and the adaptive controller.", None),
], size=17)
footer(s, 7)

# ---------------------------------------------------------------- 8. Speed results
s = slide(); bg(s)
title_bar(s, "Results", "How much faster did it get?", AMBER)
stat_card(s, 0.7, 2.0, 2.85, 1.7, "1.85x", "faster from CUDA graphs alone\n(no guessing involved)", BLUE)
stat_card(s, 3.75, 2.0, 2.85, 1.7, "2.69x", "best total speedup\n(math prompts)", GREEN)
stat_card(s, 6.8, 2.0, 2.85, 1.7, "96.7%", "of runs byte-identical\nto the original", AMBER)
stat_card(s, 9.85, 2.0, 2.85, 1.7, "0", "decoder bugs found", RED)
bullets(s, [
    ("The single biggest speed win on this hardware was NOT speculative guessing -- "
     "it was a low-level optimization (\"CUDA graphs\") that made even plain generation 1.85x faster.", None),
    ("Speculative guessing added a further, smaller boost on top of that -- but only "
     "for some kinds of prompts (math, code, structured data). It did not help, and "
     "sometimes slightly hurt, free-form writing and translation.", None),
], t=4.2, size=17)
footer(s, 8)

# ---------------------------------------------------------------- 9. The crossover point
s = slide(); bg(s)
title_bar(s, "Key Finding #1", "The \"crossover point\"", AMBER)
bullets(s, [
    ("Guessing more words per round is not always better. Each extra guess costs "
     "time whether it's right or wrong.", None),
    ("There is a point -- different for every type of prompt -- beyond which "
     "guessing MORE words per round actually makes generation SLOWER, not faster. "
     "This project calls that the crossover point.", None),
    ("Math, code and structured data can profitably guess many words ahead. Free "
     "prose and translation could not be sped up at all by guessing -- for those, "
     "plain generation was already the fastest option.", None),
    ("A surprising discovery: guessing exactly 3-4 words at once was consistently "
     "the WORST choice, across every category -- caused by a step-shaped cost "
     "curve in how the GPU verifies guesses, not a bug.", None),
], size=17)
footer(s, 9)

# ---------------------------------------------------------------- 10. The controller (honest negative)
s = slide(); bg(s)
title_bar(s, "Key Finding #2", "An honest negative result: the auto-tuner", RED)
bullets(s, [
    ("A third contribution was an adaptive controller: software that watches how "
     "often its guesses are right and automatically adjusts how many words to "
     "guess ahead, live, including turning guessing off entirely when it doesn't help.", None),
    ("Measured honestly: the adaptive controller did NOT beat simply fixing the "
     "guess count at a good constant value. Best fixed setting: 1.117x speedup. "
     "Best adaptive controller: 1.102x -- slightly worse.", None),
    ("Why: the ideal fixed setting already captures 94% of the maximum possible "
     "gain, so there was very little room left for a smart controller to win -- "
     "and its lag in reacting cost more than that small remaining room.", None),
    ("This is reported transparently as a limitation, not hidden -- it is a real "
     "and useful finding about when adaptive tuning is and isn't worth building.", None),
], size=17)
footer(s, 10)

# ---------------------------------------------------------------- 11. Architecture
s = slide(); bg(s)
title_bar(s, "Under the Hood", "System architecture", BLUE)
bullets(s, [
    ("Backend: Python, FastAPI web server, streaming results live to the browser "
     "as each word is generated (Server-Sent Events).", None),
    ("Frontend: a 4-page website built with plain HTML/CSS/JavaScript -- no "
     "external frameworks, so anyone can read and run it with zero setup steps.", None),
    ("Pages: a home page with headline results, a live interactive demo, an "
     "animated \"how it works\" walkthrough with real recorded data, and a full "
     "results dashboard with charts.", None),
    ("Every number shown in the website is read live from the actual result "
     "files on disk -- nothing on the site is typed in or made up by hand.", None),
], size=18)
footer(s, 11)

# ---------------------------------------------------------------- 12. Live demo screenshot placeholder
s = slide(); bg(s)
title_bar(s, "Live Demo", "See it in action")
tf = box(s, 0.7, 2.0, 11.9, 0.6)
para(tf, "[Insert screenshot: the live demo page, mid-generation, showing colour-coded tokens]",
     size=16, color=GREY, first=True, italic=True)
placeholder = s.shapes.add_shape(1, Inches(0.7), Inches(2.7), Inches(11.9), Inches(3.9))
placeholder.fill.solid(); placeholder.fill.fore_color.rgb = RGBColor(0xE5, 0xE7, 0xEB)
placeholder.line.color.rgb = GREY
ptf = placeholder.text_frame
p = ptf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
r = p.add_run(); r.text = "SCREENSHOT PLACEHOLDER — replace with a real capture of /demo"
r.font.size = Pt(16); r.font.color.rgb = GREY
bullets(s, [
    ("Green = draft guessed correctly. Amber = the big model corrected it. "
     "Violet = a free bonus word.", None),
], t=6.75, h=0.6, size=14)
footer(s, 12)

# ---------------------------------------------------------------- 13. Limitations
s = slide(); bg(s)
title_bar(s, "Honesty Check", "Limitations of this prototype", RED)
bullets(s, [
    ("This is a research prototype, not a production product. It was tested on "
     "ONE machine (one specific GPU, one specific Windows configuration) -- "
     "results will differ on other hardware, especially on Linux.", None),
    ("Only single-request, greedy (deterministic) generation was tested -- no "
     "batching multiple users, no creative/random sampling.", None),
    ("The evaluation set is 210 prompts, partly hand-written and partly "
     "AI-assisted and then hand-edited -- a larger, fully independent test set "
     "would strengthen the results further.", None),
    ("The adaptive controller is a validated negative result, not a working "
     "improvement -- see Key Finding #2.", None),
], size=18)
footer(s, 13)

# ---------------------------------------------------------------- 14. What's next
s = slide(); bg(s)
title_bar(s, "What's Next", "Future directions", BLUE)
bullets(s, [
    ("Test on other GPUs and on Linux, where the low-level overhead this project "
     "measured is expected to behave differently.", None),
    ("Design a smarter adaptive controller that models acceptance rate "
     "per-position rather than as one fixed number -- the identified root cause "
     "of why the current controller underperforms.", None),
    ("Extend testing to batched, multi-user serving, which is how most real "
     "products actually run these models.", None),
    ("Grow the evaluation set with more prompts and more categories for even "
     "stronger statistical confidence.", None),
], size=18)
footer(s, 14)

# ---------------------------------------------------------------- 15. Thank you / summary
s = slide(); bg(s, NAVY)
tf = box(s, 1, 1.6, 11.3, 1.0)
para(tf, "Summary", size=32, color=WHITE, bold=True, first=True, align=PP_ALIGN.CENTER)
bullets_dark = [
    "A small model guessing ahead + a big model checking in one pass = faster text, same quality",
    "Up to 2.69x measured speedup on real hardware, with zero decoder bugs found",
    "Honest science: reported what worked (fixed guessing) AND what didn't (the auto-tuner)",
    "Fully working, testable prototype with a live website demo",
]
tf2 = box(s, 1.3, 2.8, 10.7, 3.0)
for i, t in enumerate(bullets_dark):
    p = tf2.paragraphs[0] if i == 0 else tf2.add_paragraph()
    p.space_after = Pt(14)
    r = p.add_run(); r.text = "✓  " + t
    r.font.size = Pt(18); r.font.color.rgb = RGBColor(0xE5, 0xE7, 0xEB); r.font.name = "Segoe UI"
tf3 = box(s, 1, 6.3, 11.3, 0.8)
para(tf3, "Thank you  —  Project: AI Crossover Point", size=22, color=WHITE, bold=True,
     first=True, align=PP_ALIGN.CENTER)

prs.save("PRESENTATION.pptx")
print("Saved PRESENTATION.pptx with", len(prs.slides.__iter__().__length_hint__() if hasattr(prs.slides, '__length_hint__') else list(prs.slides)), "slides" if False else f"{len(prs.slides._sldIdLst)} slides")
